# -*- coding: utf-8 -*-
"""Tạo file Excel lịch dự án PharmaDash cho bài Lab QLDA (WBS, chi phí, PERT/CPM, Gantt, rủi ro)."""
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from lab_data import ACTIVITIES, RISKS, WBS, RATES, OTHER_COSTS, BUDGET, CONTINGENCY, START, DEADLINE

OUT = sys.argv[1]
F = 'Arial'
BLUE = Font(name=F, size=10, color='0000FF')
BLACK = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HDR_FONT = Font(name=F, size=10, bold=True, color='FFFFFF')
TITLE = Font(name=F, size=14, bold=True, color='1F4E79')
HDR_FILL = PatternFill('solid', fgColor='1F4E79')
L1_FILL = PatternFill('solid', fgColor='D9E2F3')
L2_FILL = PatternFill('solid', fgColor='EEF3FA')
YELLOW = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
VND = '#,##0;(#,##0);-'

wb = Workbook()


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def cell(ws, ref, value, font=BLACK, fmt=None, fill=None, align=None, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


# ============================ HƯỚNG DẪN ============================
ws = wb.active
ws.title = 'HuongDan'
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 26
ws.column_dimensions['C'].width = 90
ws['B2'] = 'LỊCH DỰ ÁN PHARMADASH – Bài Lab QLDA CNTT'
ws['B2'].font = TITLE
ws['B3'] = 'Dự án giả định: xây dựng Dashboard điều hành cho chuỗi Nhà thuốc Hữu Duyên (hư cấu). Đơn vị thời gian: ngày làm việc (T2–T6).'
ws['B3'].font = BLACK
rows = [
    ('Quy ước màu', ''),
    ('Chữ xanh dương', 'Số liệu đầu vào (giả định) – được phép sửa: giờ công, đơn giá, O/M/P, xác suất/tác động...'),
    ('Chữ đen', 'Công thức – KHÔNG sửa, tự tính lại khi đầu vào thay đổi.'),
    ('Nền vàng', 'Ô nhóm phải TỰ ĐIỀN (kết quả AI trong Claude Project để so sánh, rủi ro nhóm tự bổ sung...).'),
    ('', ''),
    ('Các sheet', ''),
    ('WBS', 'WBS 3 cấp, giờ công từng gói công việc, tự kiểm tra quy tắc 8/80 và tổng giờ (quy tắc 100%).'),
    ('ChiPhi', 'Ước lượng chi phí từ giờ công WBS × đơn giá theo vai trò + chi phí khác + dự phòng, so với ngân sách.'),
    ('PERT', 'Ước lượng 3 điểm TE = (O + 4M + P)/6, ES/EF/LS/LF, Slack, đường găng, độ lệch chuẩn, xác suất kịp hạn. '
             'Cột "TE theo AI" để nhóm nhập kết quả Claude trả lời và so sánh.'),
    ('Gantt', 'Biểu đồ Gantt tự vẽ từ ES/EF (đỏ = hoạt động găng, xanh = không găng, cột cam = hạn chót).'),
    ('NenLich', 'Kịch bản nén lịch (fast-tracking + giảm phạm vi) – lịch mới, đường găng mới, xác suất kịp hạn.'),
    ('RuiRo', 'Sổ đăng ký rủi ro (Risk Register), sắp theo điểm giảm dần. 2 dòng cuối để nhóm tự bổ sung.'),
    ('MaTran', 'Ma trận xác suất × tác động 5×5, tự điền mã rủi ro theo sheet RuiRo.'),
    ('', ''),
    ('Lưu ý', 'Các số O/M/P, giờ công, đơn giá là GIẢ ĐỊNH do AI đề xuất làm bản nháp – nhóm cần xem lại và điều chỉnh theo ý kiến của nhóm.'),
    ('', 'Nếu sửa XS/TĐ trong sheet RuiRo, hãy sắp xếp lại bảng theo cột Điểm (Data → Sort) để đúng yêu cầu "sắp theo điểm giảm dần".'),
    ('', 'Nếu thêm/bớt hoạt động PERT, phải sửa lại công thức ES (MAX EF của hoạt động trước) và LF (MIN LS của hoạt động sau).'),
]
r = 5
for a, b in rows:
    ws.cell(row=r, column=2, value=a).font = BOLD
    ws.cell(row=r, column=3, value=b).font = BLACK
    ws.cell(row=r, column=3).alignment = WRAP
    r += 1
ws['B6'].font = Font(name=F, size=10, bold=True, color='0000FF')
ws['B8'].fill = YELLOW

# ============================ WBS ============================
ws = wb.create_sheet('WBS')
ws['A1'] = 'CẤU TRÚC PHÂN RÃ CÔNG VIỆC (WBS) – PHARMADASH'
ws['A1'].font = TITLE
ws['A2'] = 'Giờ công là giả định (chữ xanh). Cấp 1–2 tự cộng từ các gói công việc cấp 3. Quy tắc 8/80: mỗi gói 8–80 giờ công.'
ws['A2'].font = BLACK
header(ws, 4, ['Mã WBS', 'Cấp', 'Tên hạng mục / gói công việc', 'Sản phẩm bàn giao', 'Vai trò chính', 'Giờ công', 'Kiểm tra 8/80'],
       [9, 6, 46, 44, 11, 10, 14])
first = 5
last = first + len(WBS) - 1
for i, (code, lvl, name, deliv, role, hours) in enumerate(WBS):
    r = first + i
    fill = L1_FILL if lvl == 1 else L2_FILL if lvl == 2 else None
    fnt = BOLD if lvl < 3 else BLACK
    cell(ws, f'A{r}', code, fnt, fill=fill)
    cell(ws, f'B{r}', lvl, fnt, fill=fill, align=CENTER)
    cell(ws, f'C{r}', ('    ' * (lvl - 1)) + name, fnt, fill=fill, align=WRAP)
    cell(ws, f'D{r}', deliv, BLACK, fill=fill, align=WRAP)
    cell(ws, f'E{r}', role, BLACK, fill=fill, align=CENTER)
    if lvl == 3:
        cell(ws, f'F{r}', hours, BLUE, '#,##0')
        cell(ws, f'G{r}', f'=IF(AND(F{r}>=8,F{r}<=80),"Đạt","Vi phạm 8/80")', BLACK, align=CENTER)
    else:
        prefix = code.split('.')[0] + '.*' if lvl == 1 else code + '.*'
        cell(ws, f'F{r}', f'=SUMIFS($F${first}:$F${last},$A${first}:$A${last},"{prefix}",$B${first}:$B${last},3)', BOLD, '#,##0', fill=fill)
        cell(ws, f'G{r}', '', fill=fill)
tr = last + 2
cell(ws, f'C{tr}', 'Tổng giờ công các gói công việc (cấp 3)', BOLD)
cell(ws, f'F{tr}', f'=SUMIFS($F${first}:$F${last},$B${first}:$B${last},3)', BOLD, '#,##0')
cell(ws, f'C{tr+1}', 'Tổng giờ công các hạng mục cấp 1 (phải bằng dòng trên – quy tắc 100%)', BOLD)
cell(ws, f'F{tr+1}', f'=SUMIFS($F${first}:$F${last},$B${first}:$B${last},1)', BOLD, '#,##0')
cell(ws, f'G{tr+1}', f'=IF(F{tr+1}=F{tr},"Khớp","Lệch")', BOLD, align=CENTER)
cell(ws, f'C{tr+2}', 'Số gói công việc vi phạm quy tắc 8/80', BOLD)
cell(ws, f'F{tr+2}', f'=COUNTIF($G${first}:$G${last},"Vi phạm 8/80")', BOLD, '0')
ws.freeze_panes = 'A5'
WBS_TOTAL = f'WBS!$F${tr}'
WBS_RANGE = (first, last)

# ============================ CHI PHÍ ============================
ws = wb.create_sheet('ChiPhi')
ws['A1'] = 'ƯỚC LƯỢNG CHI PHÍ DỰ ÁN (VND)'
ws['A1'].font = TITLE
ws['A2'] = 'Đơn giá là giả định của nhóm (chi phí công ty gồm lương + bảo hiểm, tính theo giờ), không phải báo giá thật.'
ws['A2'].font = BLACK
header(ws, 4, ['Vai trò', 'Diễn giải', 'Đơn giá (VND/giờ)', 'Giờ công (từ WBS)', 'Thành tiền (VND)'], [14, 38, 18, 18, 20])
fr, lr = WBS_RANGE
r = 5
for role, desc, rate in RATES:
    cell(ws, f'A{r}', role, BOLD)
    cell(ws, f'B{r}', desc)
    cell(ws, f'C{r}', rate, BLUE, VND)
    cell(ws, f'D{r}', f'=SUMIFS(WBS!$F${fr}:$F${lr},WBS!$E${fr}:$E${lr},A{r},WBS!$B${fr}:$B${lr},3)', Font(name=F, size=10, color='008000'), '#,##0')
    cell(ws, f'E{r}', f'=C{r}*D{r}', BLACK, VND)
    r += 1
lab_last = r - 1
cell(ws, f'A{r}', 'Cộng nhân công', BOLD)
cell(ws, f'D{r}', f'=SUM(D5:D{lab_last})', BOLD, '#,##0')
cell(ws, f'E{r}', f'=SUM(E5:E{lab_last})', BOLD, VND)
labor = r
cell(ws, f'B{r+1}', 'Kiểm tra: giờ công khớp tổng WBS?', BLACK)
cell(ws, f'D{r+1}', f'=IF(D{r}={WBS_TOTAL},"Khớp","Lệch")', BLACK, align=CENTER)
r += 3
header_row = r
for i, t in enumerate(['Chi phí khác', 'Diễn giải', 'Số lượng', 'Đơn giá (VND)', 'Thành tiền (VND)'], 1):
    c = ws.cell(row=r, column=i, value=t)
    c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
r += 1
o_first = r
for name, desc, qty, price in OTHER_COSTS:
    cell(ws, f'A{r}', name, BOLD)
    cell(ws, f'B{r}', desc, align=WRAP)
    cell(ws, f'C{r}', qty, BLUE, '#,##0')
    cell(ws, f'D{r}', price, BLUE, VND)
    cell(ws, f'E{r}', f'=C{r}*D{r}', BLACK, VND)
    r += 1
o_last = r - 1
cell(ws, f'A{r}', 'Cộng chi phí khác', BOLD)
cell(ws, f'E{r}', f'=SUM(E{o_first}:E{o_last})', BOLD, VND)
other = r
r += 2
cell(ws, f'A{r}', 'Chi phí trực tiếp', BOLD); cell(ws, f'E{r}', f'=E{labor}+E{other}', BOLD, VND); direct = r; r += 1
cell(ws, f'A{r}', 'Tỷ lệ dự phòng', BOLD); cell(ws, f'B{r}', 'Dự phòng rủi ro (contingency reserve)'); cell(ws, f'C{r}', CONTINGENCY, BLUE, '0%'); cell(ws, f'E{r}', f'=E{direct}*C{r}', BLACK, VND); cont = r; r += 1
cell(ws, f'A{r}', 'TỔNG DỰ TOÁN', BOLD); cell(ws, f'E{r}', f'=E{direct}+E{cont}', BOLD, VND); total = r; r += 1
cell(ws, f'A{r}', 'Ngân sách được duyệt', BOLD); cell(ws, f'B{r}', 'Theo tài liệu yêu cầu khách hàng'); cell(ws, f'E{r}', BUDGET, BLUE, VND); bud = r; r += 1
cell(ws, f'A{r}', 'Còn lại', BOLD); cell(ws, f'E{r}', f'=E{bud}-E{total}', BOLD, VND); r += 1
cell(ws, f'A{r}', 'Tỷ lệ sử dụng', BOLD); cell(ws, f'E{r}', f'=E{total}/E{bud}', BOLD, '0.0%'); r += 1
cell(ws, f'A{r}', 'Kết luận', BOLD); cell(ws, f'E{r}', f'=IF(E{total}<=E{bud},"Trong ngân sách","VƯỢT NGÂN SÁCH")', BOLD, align=CENTER)

# ============================ PERT ============================
ws = wb.create_sheet('PERT')
ws['A1'] = 'ƯỚC LƯỢNG PERT 3 ĐIỂM & ĐƯỜNG GĂNG (CPM) – đơn vị: ngày làm việc'
ws['A1'].font = TITLE
cell(ws, 'A3', 'Ngày khởi động', BOLD); cell(ws, 'C3', START, BLUE, 'dd/mm/yyyy')
cell(ws, 'A4', 'Hạn chót go-live', BOLD); cell(ws, 'C4', DEADLINE, BLUE, 'dd/mm/yyyy')
cell(ws, 'A5', 'Số ngày làm việc tới hạn', BOLD); cell(ws, 'C5', '=NETWORKDAYS(C3,C4)', BOLD, '0')
ws['D3'] = 'Ngày làm việc thứ 1 = ngày khởi động. Chưa trừ ngày nghỉ lễ (giai đoạn 10–12/2026 không có ngày lễ).'
ws['D3'].font = BLACK
cols = ['Mã', 'Hoạt động', 'WBS liên quan', 'Hoạt động trước', 'O', 'M', 'P', 'TE = (O+4M+P)/6', 'σ = (P−O)/6', 'σ²',
        'ES', 'EF', 'LS', 'LF', 'Slack', 'Găng?', 'Ngày bắt đầu', 'Ngày kết thúc', 'TE theo AI (nhóm nhập)', 'Chênh lệch Excel − AI']
header(ws, 8, cols, [6, 38, 14, 11, 6, 6, 6, 11, 9, 8, 8, 8, 8, 8, 8, 8, 12, 12, 14, 14])
ws.row_dimensions[8].height = 42
row_of = {}
f0 = 9
for i, a in enumerate(ACTIVITIES):
    row_of[a['id']] = f0 + i
lastp = f0 + len(ACTIVITIES) - 1
succ = {a['id']: [] for a in ACTIVITIES}
for a in ACTIVITIES:
    for p in a['pred']:
        succ[p].append(a['id'])
T_CELL = '$C$' + str(lastp + 3)
for a in ACTIVITIES:
    r = row_of[a['id']]
    cell(ws, f'A{r}', a['id'], BOLD, align=CENTER)
    cell(ws, f'B{r}', a['name'], align=WRAP)
    cell(ws, f'C{r}', a['wbs'], align=CENTER)
    cell(ws, f'D{r}', ', '.join(a['pred']) or '—', align=CENTER)
    cell(ws, f'E{r}', a['o'], BLUE, '0')
    cell(ws, f'F{r}', a['m'], BLUE, '0')
    cell(ws, f'G{r}', a['p'], BLUE, '0')
    cell(ws, f'H{r}', f'=(E{r}+4*F{r}+G{r})/6', BLACK, '0.00')
    cell(ws, f'I{r}', f'=(G{r}-E{r})/6', BLACK, '0.00')
    cell(ws, f'J{r}', f'=I{r}^2', BLACK, '0.00')
    es = '0' if not a['pred'] else 'MAX(' + ','.join(f'L{row_of[p]}' for p in a['pred']) + ')'
    cell(ws, f'K{r}', '=' + es, BLACK, '0.00')
    cell(ws, f'L{r}', f'=K{r}+H{r}', BLACK, '0.00')
    cell(ws, f'M{r}', f'=N{r}-H{r}', BLACK, '0.00')
    lf = T_CELL if not succ[a['id']] else 'MIN(' + ','.join(f'M{row_of[s]}' for s in succ[a['id']]) + ')'
    cell(ws, f'N{r}', '=' + lf, BLACK, '0.00')
    cell(ws, f'O{r}', f'=M{r}-K{r}', BLACK, '0.00')
    cell(ws, f'P{r}', f'=IF(ABS(O{r})<0.005,"Găng","")', BOLD, align=CENTER)
    cell(ws, f'Q{r}', f'=WORKDAY($C$3,INT(K{r}))', BLACK, 'dd/mm/yyyy')
    cell(ws, f'R{r}', f'=WORKDAY($C$3,ROUNDUP(L{r},0)-1)', BLACK, 'dd/mm/yyyy')
    cell(ws, f'S{r}', None, BLUE, '0.00', fill=YELLOW)
    cell(ws, f'T{r}', f'=IF(S{r}="","",H{r}-S{r})', BLACK, '0.00')
ws.conditional_formatting.add(f'A{f0}:P{lastp}', FormulaRule(formula=[f'$P{f0}="Găng"'], fill=PatternFill('solid', start_color='FCE4E4', end_color='FCE4E4')))
s = lastp + 2
ws[f'A{s}'] = 'KẾT QUẢ'
ws[f'A{s}'].font = TITLE
items = [
    ('Thời gian dự án kỳ vọng T (ngày làm việc)', f'=MAX(L{f0}:L{lastp})', '0.00'),
    ('Đường găng', '=' + '&'.join(f'IF(P{row_of[a["id"]]}="Găng","{a["id"]} → ","")' for a in ACTIVITIES), None),
    ('Tổng phương sai đường găng Σσ²', f'=SUMIF(P{f0}:P{lastp},"Găng",J{f0}:J{lastp})', '0.00'),
    ('Độ lệch chuẩn đường găng σ', f'=SQRT(C{s+3})', '0.00'),
    ('Số ngày làm việc tới hạn chót', '=C5', '0'),
    ('Chênh lệch T − hạn chót (ngày; dương = trễ)', f'=C{s+1}-C{s+5}', '0.00'),
    ('Z = (hạn chót − T) / σ', f'=(C{s+5}-C{s+1})/C{s+4}', '0.00'),
    ('Xác suất hoàn thành đúng hạn P(Z)', f'=NORMSDIST(C{s+7})', '0.00%'),
    ('Ngày kết thúc dự kiến (T)', f'=WORKDAY(C3,ROUNDUP(C{s+1},0)-1)', 'dd/mm/yyyy'),
    ('Kết luận', f'=IF(C{s+6}<=0,"Kịp hạn chót","TRỄ HẠN – cần nén lịch (crashing / fast-tracking)")', None),
]
for i, (label, formula, fmt) in enumerate(items, 1):
    rr = s + i
    ws.merge_cells(f'A{rr}:B{rr}')
    cell(ws, f'A{rr}', label, BOLD)
    ws[f'B{rr}'].border = BORDER
    cell(ws, f'C{rr}', formula, BOLD, fmt)
    if label == 'Đường găng':
        # bỏ mũi tên thừa ở cuối chuỗi
        cell(ws, f'C{rr}', f'=LEFT({formula[1:]},LEN({formula[1:]})-3)', BOLD)
        ws.merge_cells(f'C{rr}:J{rr}')
PERT_T = f'PERT!$C${s+1}'
ws.freeze_panes = 'C9'

# ============================ GANTT ============================
ws = wb.create_sheet('Gantt')
ws['A1'] = 'BIỂU ĐỒ GANTT (tự vẽ từ sheet PERT)'
ws['A1'].font = TITLE
ws['A2'] = 'Đỏ = hoạt động găng · Xanh = không găng · Cột cam = ngày làm việc cuối cùng trước hạn chót. Mỗi cột = 1 ngày làm việc.'
ws['A2'].font = BLACK
NDAYS = 85
header(ws, 5, ['Mã', 'Hoạt động', 'ES', 'EF', 'Găng?'], [5, 46, 6, 6, 7])
for d in range(1, NDAYS + 1):
    col = 5 + d
    c4 = ws.cell(row=4, column=col, value=f'=WORKDAY(PERT!$C$3,{d - 1})')
    c4.number_format = 'dd/mm'
    c4.font = Font(name=F, size=7)
    c4.alignment = Alignment(text_rotation=90, horizontal='center')
    c5 = ws.cell(row=5, column=col, value=d)
    c5.font, c5.fill, c5.alignment, c5.border = Font(name=F, size=7, bold=True, color='FFFFFF'), HDR_FILL, CENTER, BORDER
    ws.column_dimensions[L(col)].width = 2.6
ws.row_dimensions[4].height = 34
for i, a in enumerate(ACTIVITIES):
    r = 6 + i
    pr = row_of[a['id']]
    cell(ws, f'A{r}', f'=PERT!A{pr}', BOLD, align=CENTER)
    cell(ws, f'B{r}', f'=PERT!B{pr}')
    cell(ws, f'C{r}', f'=PERT!K{pr}', BLACK, '0.0')
    cell(ws, f'D{r}', f'=PERT!L{pr}', BLACK, '0.0')
    cell(ws, f'E{r}', f'=PERT!P{pr}', BOLD, align=CENTER)
    for d in range(1, NDAYS + 1):
        col = L(5 + d)
        c = ws[f'{col}{r}']
        c.value = f'=IF(AND({col}$5>$C{r},{col}$5-1<$D{r}),1,"")'
        c.number_format = ';;;'
        c.border = Border(left=Side(style='hair', color='E0E0E0'), right=Side(style='hair', color='E0E0E0'),
                          top=thin, bottom=thin)
g_last = 6 + len(ACTIVITIES) - 1
rng = f'F6:{L(5 + NDAYS)}{g_last}'
ws.conditional_formatting.add(rng, FormulaRule(formula=['AND(F6=1,$E6="Găng")'], fill=PatternFill('solid', start_color='E03C3C', end_color='E03C3C')))
ws.conditional_formatting.add(rng, FormulaRule(formula=['F6=1'], fill=PatternFill('solid', start_color='5B8DEF', end_color='5B8DEF')))
ws.conditional_formatting.add(f'F4:{L(5 + NDAYS)}5', FormulaRule(formula=['F$5=PERT!$C$5'], fill=PatternFill('solid', start_color='F5A623', end_color='F5A623')))
ws.conditional_formatting.add(rng, FormulaRule(formula=['F$5=PERT!$C$5'], fill=PatternFill('solid', start_color='FDE7C2', end_color='FDE7C2')))
ws.freeze_panes = 'F6'

# ============================ NÉN LỊCH (kịch bản) ============================
ws = wb.create_sheet('NenLich')
ws['A1'] = 'PHƯƠNG ÁN NÉN LỊCH – fast-tracking + giảm phạm vi (AI đề xuất, nhóm xác nhận)'
ws['A1'].font = TITLE
ws['A2'] = ('Cột "Chồng lấn" = số ngày hoạt động được bắt đầu SỚM trước khi các hoạt động trước kết thúc (fast-tracking). '
            'ES = MAX(0; MAX(EF hoạt động trước) − Chồng lấn); LF = MIN(LS hoạt động sau + Chồng lấn của hoạt động sau). Sheet PERT giữ nguyên làm lịch gốc.')
ws['A2'].font = BLACK
ws['A2'].alignment = WRAP
ws.merge_cells('A2:P2')
ws.row_dimensions[2].height = 30
cols = ['Mã', 'Hoạt động', 'Trước', 'O', 'M', 'P', 'Chồng lấn', 'TE', 'σ²', 'ES', 'EF', 'LS', 'LF', 'Slack', 'Găng?', 'Thay đổi so với lịch gốc']
header(ws, 4, cols, [6, 36, 12, 6, 6, 6, 9, 8, 8, 8, 8, 8, 8, 8, 8, 70])
ws.row_dimensions[4].height = 30
from lab_data import SCENARIO
n0 = 5
nrow = {a['id']: n0 + i for i, a in enumerate(ACTIVITIES)}
nlast = n0 + len(ACTIVITIES) - 1
NT = '$C$' + str(nlast + 3)
for a in ACTIVITIES:
    r = nrow[a['id']]
    sc = SCENARIO.get(a['id'], {})
    changed = PatternFill('solid', fgColor='FFF2CC')
    cell(ws, f'A{r}', a['id'], BOLD, align=CENTER)
    cell(ws, f'B{r}', f'=PERT!B{row_of[a["id"]]}', align=WRAP)
    cell(ws, f'C{r}', ', '.join(a['pred']) or '—', align=CENTER)
    for col, key in (('D', 'o'), ('E', 'm'), ('F', 'p')):
        if key in sc:
            cell(ws, f'{col}{r}', sc[key], BLUE, '0', fill=changed)
        else:
            cell(ws, f'{col}{r}', f'=PERT!{ {"o": "E", "m": "F", "p": "G"}[key] }{row_of[a["id"]]}', Font(name=F, size=10, color='008000'), '0')
    cell(ws, f'G{r}', sc.get('ov', 0), BLUE, '0.0', fill=changed if 'ov' in sc else None)
    cell(ws, f'H{r}', f'=(D{r}+4*E{r}+F{r})/6', BLACK, '0.00')
    cell(ws, f'I{r}', f'=((F{r}-D{r})/6)^2', BLACK, '0.00')
    es = '0' if not a['pred'] else f'MAX(0,MAX(' + ','.join(f'K{nrow[p]}' for p in a['pred']) + f')-G{r})'
    cell(ws, f'J{r}', '=' + es, BLACK, '0.00')
    cell(ws, f'K{r}', f'=J{r}+H{r}', BLACK, '0.00')
    cell(ws, f'L{r}', f'=M{r}-H{r}', BLACK, '0.00')
    lf = NT if not succ[a['id']] else 'MIN(' + ','.join(f'L{nrow[s]}+G{nrow[s]}' for s in succ[a['id']]) + ')'
    cell(ws, f'M{r}', '=' + lf, BLACK, '0.00')
    cell(ws, f'N{r}', f'=L{r}-J{r}', BLACK, '0.00')
    cell(ws, f'O{r}', f'=IF(ABS(N{r})<0.005,"Găng","")', BOLD, align=CENTER)
    cell(ws, f'P{r}', sc.get('note', ''), align=WRAP)
ws.conditional_formatting.add(f'A{n0}:O{nlast}', FormulaRule(formula=[f'$O{n0}="Găng"'], fill=PatternFill('solid', start_color='FCE4E4', end_color='FCE4E4')))
s2 = nlast + 2
ws[f'A{s2}'] = 'KẾT QUẢ SAU KHI NÉN'
ws[f'A{s2}'].font = TITLE
res = [
    ('Thời gian dự án T mới (ngày làm việc)', f'=MAX(K{n0}:K{nlast})', '0.00'),
    ('Đường găng mới', None, None),
    ('Độ lệch chuẩn đường găng σ', f'=SQRT(SUMIF(O{n0}:O{nlast},"Găng",I{n0}:I{nlast}))', '0.00'),
    ('Số ngày làm việc tới hạn chót', '=PERT!C5', '0'),
    ('Dự trữ so với hạn chót (ngày)', f'=C{s2+4}-C{s2+1}', '0.00'),
    ('Xác suất kịp hạn P(Z)', f'=NORMSDIST((C{s2+4}-C{s2+1})/C{s2+3})', '0.0%'),
    ('Ngày kết thúc dự kiến', f'=WORKDAY(PERT!C3,ROUNDUP(C{s2+1},0)-1)', 'dd/mm/yyyy'),
    ('T lịch gốc (sheet PERT) / số ngày rút ngắn', f'={PERT_T}', '0.00'),
]
for i, (label, formula, fmt) in enumerate(res, 1):
    rr = s2 + i
    ws.merge_cells(f'A{rr}:B{rr}')
    cell(ws, f'A{rr}', label, BOLD)
    ws[f'B{rr}'].border = BORDER
    if label == 'Đường găng mới':
        expr = '&'.join(f'IF(O{nrow[a["id"]]}="Găng","{a["id"]} → ","")' for a in ACTIVITIES)
        cell(ws, f'C{rr}', f'=LEFT({expr},LEN({expr})-3)', BOLD)
        ws.merge_cells(f'C{rr}:J{rr}')
    else:
        cell(ws, f'C{rr}', formula, BOLD, fmt)
cell(ws, f'D{s2+8}', f'=C{s2+8}-C{s2+1}', BOLD, '0.00')
ws.freeze_panes = 'C5'

# ============================ RỦI RO ============================
ws = wb.create_sheet('RuiRo')
ws['A1'] = 'SỔ ĐĂNG KÝ RỦI RO (RISK REGISTER) – sắp theo điểm giảm dần'
ws['A1'].font = TITLE
ws['A2'] = 'Mô tả theo cấu trúc Nguyên nhân – Sự kiện – Hậu quả. XS, TĐ thang 1–5; Điểm = XS × TĐ; ≥15 Cao, 8–14 Trung bình, <8 Thấp.'
ws['A2'].font = BLACK
header(ws, 4, ['ID', 'Nguyên nhân', 'Sự kiện rủi ro', 'Hậu quả', 'Nhóm', 'XS (1–5)', 'TĐ (1–5)', 'Điểm', 'Mức',
               'Chiến lược', 'Hành động ứng phó', 'Người phụ trách', 'Nguồn'],
       [6, 30, 30, 30, 13, 8, 8, 7, 11, 11, 44, 13, 12])
ws.row_dimensions[4].height = 30
risks = sorted(RISKS, key=lambda x: -x['xs'] * x['td'])
r = 5
for k in risks:
    vals = [k['id'], k['cause'], k['event'], k['impact'], k['group']]
    for j, v in enumerate(vals, 1):
        cell(ws, f'{L(j)}{r}', v, BOLD if j == 1 else BLACK, align=WRAP)
    cell(ws, f'F{r}', k['xs'], BLUE, '0', align=CENTER)
    cell(ws, f'G{r}', k['td'], BLUE, '0', align=CENTER)
    cell(ws, f'H{r}', f'=F{r}*G{r}', BOLD, '0', align=CENTER)
    cell(ws, f'I{r}', f'=IF(H{r}>=15,"Cao",IF(H{r}>=8,"Trung bình","Thấp"))', BOLD, align=CENTER)
    cell(ws, f'J{r}', k['strategy'], align=WRAP)
    cell(ws, f'K{r}', k['action'], align=WRAP)
    cell(ws, f'L{r}', k['owner'], align=WRAP)
    cell(ws, f'M{r}', 'AI gợi ý', align=CENTER)
    r += 1
for extra in ('R11', 'R12'):
    cell(ws, f'A{r}', extra, BOLD)
    for j in range(2, 14):
        c = ws.cell(row=r, column=j)
        c.fill, c.border, c.font, c.alignment = YELLOW, BORDER, BLUE, WRAP
    ws[f'H{r}'] = f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})'
    ws[f'I{r}'] = f'=IF(H{r}="","",IF(H{r}>=15,"Cao",IF(H{r}>=8,"Trung bình","Thấp")))'
    ws[f'M{r}'] = 'Nhóm bổ sung'
    ws.row_dimensions[r].height = 40
    r += 1
risk_last = r - 1
ws.conditional_formatting.add(f'I5:I{risk_last}', FormulaRule(formula=['I5="Cao"'], fill=PatternFill('solid', start_color='F8CBCB', end_color='F8CBCB')))
ws.conditional_formatting.add(f'I5:I{risk_last}', FormulaRule(formula=['I5="Trung bình"'], fill=PatternFill('solid', start_color='FCE4C4', end_color='FCE4C4')))
ws.conditional_formatting.add(f'I5:I{risk_last}', FormulaRule(formula=['I5="Thấp"'], fill=PatternFill('solid', start_color='D8F0DD', end_color='D8F0DD')))
ws.freeze_panes = 'B5'

# ============================ MA TRẬN ============================
ws = wb.create_sheet('MaTran')
ws['A1'] = 'MA TRẬN XÁC SUẤT × TÁC ĐỘNG (5×5) – tự cập nhật từ sheet RuiRo'
ws['A1'].font = TITLE
ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 8
for c in 'CDEFG':
    ws.column_dimensions[c].width = 18
ws['C3'] = 'TÁC ĐỘNG →'
ws['C3'].font = BOLD
ws.merge_cells('C3:G3')
ws['C3'].alignment = CENTER
for j, td in enumerate(range(1, 6)):
    c = ws.cell(row=4, column=3 + j, value=td)
    c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
ws['A5'] = 'XÁC SUẤT ↓'
ws['A5'].font = BOLD
for i, xs in enumerate(range(5, 0, -1)):
    r = 5 + i
    ws.row_dimensions[r].height = 42
    c = ws.cell(row=r, column=2, value=xs)
    c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    for j, td in enumerate(range(1, 6)):
        parts = [f'IF(AND(RuiRo!$F${k}={xs},RuiRo!$G${k}={td}),RuiRo!$A${k}&" ","")' for k in range(5, risk_last + 1)]
        c = ws.cell(row=r, column=3 + j, value='=TRIM(' + '&'.join(parts) + ')')
        score = xs * td
        color = 'F8CBCB' if score >= 15 else 'FCE4C4' if score >= 8 else 'D8F0DD'
        c.fill = PatternFill('solid', fgColor=color)
        c.font, c.alignment, c.border = BOLD, CENTER, BORDER
ws['A11'] = 'Vùng đỏ: điểm ≥ 15 (Cao) · Cam: 8–14 (Trung bình) · Xanh: < 8 (Thấp)'
ws['A11'].font = BLACK

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = sh.title in ('Gantt',)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved')
